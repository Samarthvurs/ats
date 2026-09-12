from fastapi import FastAPI, UploadFile, File, Form
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel
from typing import List, Optional
import pdfplumber
import docx as docx_lib
import io
import json
import time
import uuid
import threading
import traceback
from services.nlp_engine import rank_resume
from services.shortlist_engine import rank_candidates, extract_skills_from_jd, detect_bias_flags

app = FastAPI(title="Resumeit ATS ML Engine")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

class ScoreRequest(BaseModel):
    text: str
    job_description: str = ""

@app.get("/")
def root():
    return {"message": "Resumeeit ML API is Online"}

@app.get("/health")
def health():
    return {"status": "ok", "message": "ML ATS Engine is running natively!"}

@app.post("/api/score-resume")
async def score_resume(request: ScoreRequest):
    try:
        result = rank_resume(request.text, request.job_description)
        return {"success": True, "data": result}
    except Exception as e:
        print(traceback.format_exc())
        return {"success": False, "error": str(e)}

@app.post("/api/parse-pdf")
async def parse_pdf(file: UploadFile = File(...)):
    text = ""
    try:
        content = await file.read()
        with pdfplumber.open(io.BytesIO(content)) as pdf:
            for page in pdf.pages:
                extracted = page.extract_text()
                if extracted:
                    text += extracted + "\n"

        print(f"DEBUG: Extracted {len(text)} chars from PDF. Preview: {text[:200]}...")
        return {"success": True, "text": text}
    except Exception as e:
        print(traceback.format_exc())
        return {"success": False, "error": str(e)}


def extract_text_from_upload(filename: str, content: bytes) -> str:
    """Generalized text extraction — supports PDF and DOCX, used for both
    the JD and every resume in a batch."""
    ext = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""

    if ext == "pdf":
        text = ""
        with pdfplumber.open(io.BytesIO(content)) as pdf:
            for page in pdf.pages:
                extracted = page.extract_text()
                if extracted:
                    text += extracted + "\n"
        return text

    if ext == "docx":
        document = docx_lib.Document(io.BytesIO(content))
        parts = [p.text for p in document.paragraphs if p.text]
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text:
                        parts.append(cell.text)
        return "\n".join(parts)

    # Fallback: plain text upload
    try:
        return content.decode("utf-8", errors="ignore")
    except Exception:
        return ""


@app.post("/api/parse-document")
async def parse_document(file: UploadFile = File(...)):
    """Parses a JD (or any single document) — PDF or DOCX — to plain text."""
    try:
        content = await file.read()
        text = extract_text_from_upload(file.filename, content)
        return {"success": True, "text": text}
    except Exception as e:
        print(traceback.format_exc())
        return {"success": False, "error": str(e)}


class DetectSkillsRequest(BaseModel):
    job_description: str


@app.post("/api/detect-skills")
async def detect_skills(request: DetectSkillsRequest):
    """Lightweight preview of what the matching engine will treat as
    required/preferred skills — used by the JD step so the recruiter can see
    (and add to) detected skills before uploading any resumes."""
    try:
        skills = extract_skills_from_jd(request.job_description)
        bias_flags = detect_bias_flags(request.job_description)
        return {"success": True, "data": {**skills, "bias_flags": bias_flags}}
    except Exception as e:
        print(traceback.format_exc())
        return {"success": False, "error": str(e)}


# =========================================================================
# Smart Shortlisting Engine — async job pipeline
# =========================================================================
# 200+ resumes can take a while to parse + embed, so shortlisting runs as a
# background job: the client gets a job_id immediately and polls for
# progress instead of holding one long HTTP request open.
# =========================================================================

JOBS = {}
JOBS_LOCK = threading.Lock()
JOB_TTL_SECONDS = 60 * 60  # jobs are dropped an hour after completion


def _set_job(job_id: str, **fields):
    with JOBS_LOCK:
        if job_id in JOBS:
            JOBS[job_id].update(fields)


def _process_job(job_id: str, jd_text: str, resume_uploads: list, config: dict):
    try:
        total = len(resume_uploads)
        _set_job(job_id, status="parsing", progress={"processed": 0, "total": total, "phase": "parsing"})

        parsed = []
        for i, (filename, content) in enumerate(resume_uploads):
            entry = {"filename": filename, "text": "", "parse_error": None}
            try:
                text = extract_text_from_upload(filename, content)
                if len(text.strip()) < 40:
                    entry["parse_error"] = "No readable text extracted (scanned image, empty, or unsupported format)."
                entry["text"] = text
            except Exception as e:
                entry["parse_error"] = f"Failed to parse file: {e}"
            parsed.append(entry)
            if i % 5 == 0 or i == total - 1:
                _set_job(job_id, progress={"processed": i + 1, "total": total, "phase": "parsing"})

        _set_job(job_id, status="matching", progress={"processed": 0, "total": total, "phase": "embedding_skills"})

        def phase_cb(phase_name):
            _set_job(job_id, progress={"processed": 0, "total": total, "phase": phase_name})

        def progress_cb(done, total_n):
            if done % 3 == 0 or done == total_n:
                _set_job(job_id, progress={"processed": done, "total": total_n, "phase": "scoring_candidates"})

        result = rank_candidates(jd_text, parsed, config=config, progress_cb=progress_cb, phase_cb=phase_cb)

        _set_job(
            job_id,
            status="done",
            result=result,
            progress={"processed": total, "total": total, "phase": "done"},
            finished_at=time.time(),
        )
    except Exception as e:
        print(traceback.format_exc())
        _set_job(job_id, status="error", error=str(e), finished_at=time.time())


def _reap_old_jobs():
    now = time.time()
    with JOBS_LOCK:
        stale = [jid for jid, j in JOBS.items() if j.get("finished_at") and now - j["finished_at"] > JOB_TTL_SECONDS]
        for jid in stale:
            del JOBS[jid]


@app.post("/api/shortlist/jobs")
async def create_shortlist_job(
    job_description: str = Form(""),
    jd_file: Optional[UploadFile] = File(None),
    resumes: List[UploadFile] = File(...),
    config: str = Form("{}"),
):
    """
    Starts a shortlisting job for a batch of resumes (handles hundreds).
    Returns a job_id immediately — poll GET /api/shortlist/jobs/{job_id}.
    """
    _reap_old_jobs()

    jd_text = (job_description or "").strip()
    if jd_file is not None and jd_file.filename:
        try:
            content = await jd_file.read()
            parsed_jd = extract_text_from_upload(jd_file.filename, content)
            if parsed_jd.strip():
                jd_text = parsed_jd
        except Exception as e:
            return {"success": False, "error": f"Could not parse JD file: {e}"}

    if not jd_text.strip():
        return {"success": False, "error": "No job description text or file provided."}
    if not resumes:
        return {"success": False, "error": "No resumes uploaded."}

    try:
        cfg = json.loads(config) if config else {}
    except Exception:
        cfg = {}

    resume_uploads = []
    for f in resumes:
        content = await f.read()
        resume_uploads.append((f.filename, content))

    job_id = str(uuid.uuid4())
    with JOBS_LOCK:
        JOBS[job_id] = {
            "status": "queued",
            "progress": {"processed": 0, "total": len(resume_uploads), "phase": "queued"},
            "created_at": time.time(),
        }

    thread = threading.Thread(target=_process_job, args=(job_id, jd_text, resume_uploads, cfg), daemon=True)
    thread.start()

    return {"success": True, "job_id": job_id}


@app.get("/api/shortlist/jobs/{job_id}")
async def get_shortlist_job(job_id: str):
    with JOBS_LOCK:
        job = JOBS.get(job_id)
    if not job:
        return {"success": False, "error": "Job not found or expired."}
    return {"success": True, "data": job}
